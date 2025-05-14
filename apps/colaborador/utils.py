import random
import string
import datetime
from builtins import range
from django.contrib.admin.models import ADDITION, LogEntry
from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import ValidationError
from django.utils.encoding import force_str, force_text
from django.utils.http import urlsafe_base64_decode
from apps.colaborador.models import Colaborador
import pandas as pd
from io import BytesIO
from django.http import HttpResponse
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
import os
from django.conf import settings
from PIL import Image as PILImage

def gerar_password():
    password = ""
    for ay in list(range(3)):
        i = random.choice(string.ascii_lowercase)
        password += i
    password += random.choice(string.punctuation )
    for ay in list(range(2)):
        o = random.choice(string.digits)
        password += o
    password += random.choice(string.punctuation )
    for ay in list(range(3)):
        u = random.choice(string.ascii_uppercase)
        password += u
    return password


def get_user(uidb64):
    try:
        pk = urlsafe_base64_decode(uidb64).decode()
        user = Colaborador._default_manager.get(pk=pk)
    except (TypeError, ValueError, OverflowError, Colaborador.DoesNotExist, ValidationError):
        user = None
    return user


class HistoryColaborador:
    def __init__(self, request=None):
        self.request = request

    def responsavel(self, colaborador_grupoacesso, status):
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador_grupoacesso.colaborador).pk, object_id=colaborador_grupoacesso.colaborador.pk,
            object_repr=force_str(colaborador_grupoacesso.colaborador), action_flag=ADDITION,
            change_message=(f"Solicitação de acesso aos recursos do Grupo de Trabalho {colaborador_grupoacesso.grupo_acesso.grupo_acesso} - {status}")
        )
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador_grupoacesso.grupo_acesso).pk,
            object_id=colaborador_grupoacesso.grupo_acesso.pk,
            object_repr=force_str(colaborador_grupoacesso.grupo_acesso), action_flag=ADDITION,
            change_message=(f"Solicitação do Colaborador {colaborador_grupoacesso.colaborador.full_name} - {status}"),
        )
    
    def responsavel_remover_grupo_acesso(self, colaborador_grupoacesso):
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador_grupoacesso.colaborador).pk, object_id=colaborador_grupoacesso.colaborador.pk,
            object_repr=force_str(colaborador_grupoacesso.colaborador), action_flag=ADDITION,
            change_message=(f"O acesso aos recursos do Grupo de Trabalho {colaborador_grupoacesso.grupo_acesso.grupo_acesso} foram removidos")
        )
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador_grupoacesso.grupo_acesso).pk,
            object_id=colaborador_grupoacesso.grupo_acesso.pk,
            object_repr=force_str(colaborador_grupoacesso.grupo_acesso), action_flag=ADDITION,
            change_message=(f"O acesso do Colaborador {colaborador_grupoacesso.colaborador} foi removido")
        )
    
    def responsavel_remover_grupo(self, colaborador_grupoacesso):
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador_grupoacesso.colaborador).pk, object_id=colaborador_grupoacesso.colaborador.pk,
            object_repr=force_str(colaborador_grupoacesso.colaborador), action_flag=ADDITION,
            change_message=(f"O acesso ao Grupo de Trabalho {colaborador_grupoacesso.grupo_acesso.grupo_acesso} foi removido")
        )
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador_grupoacesso.grupo_acesso.grupo_trabalho).pk,
            object_id=colaborador_grupoacesso.grupo_acesso.grupo_trabalho.pk,
            object_repr=force_str(colaborador_grupoacesso.grupo_acesso.grupo_trabalho), action_flag=ADDITION,
            change_message=(f"O acesso do Colaborador {colaborador_grupoacesso.colaborador} foi removido")
        )

    def novo(self, colaborador):
        LogEntry.objects.log_action( user_id=colaborador.pk, 
        content_type_id=ContentType.objects.get_for_model(colaborador).pk, 
        object_id=colaborador.pk, 
        object_repr=force_str(colaborador), action_flag=ADDITION, 
        change_message=("Solicitação de Cadastro")
        )

    def suporte(self, colaborador):
        LogEntry.objects.log_action(
            user_id=self.request.user.pk, content_type_id=ContentType.objects.get_for_model(colaborador).pk, object_id=colaborador.pk, object_repr=force_str(colaborador), action_flag=ADDITION, change_message=("Conta criada pelo suporte")
        )

    def secretaria(self, colaborador):
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador).pk,
            object_id=colaborador.pk,
            object_repr=force_str(colaborador), action_flag=ADDITION,
            change_message=(f"Cadastro revisado pela Secretaria da {colaborador.divisao.divisao}"),
        )

    def chefia(self, colaborador):
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador).pk,
            object_id=colaborador.pk,
            object_repr=force_str(colaborador), action_flag=ADDITION,
            change_message=(f"Cadastro aprovado pela chefia da {colaborador.divisao.divisao}"),
        )

    def solicitacao(self, colaborador, grupo_acesso):
        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(colaborador).pk,
            object_id=colaborador.pk,
            object_repr=force_str(colaborador),
            action_flag=ADDITION,
            change_message=(f"Solicitação de acesso aos recursos do Grupo de Trabalho {grupo_acesso.grupo_acesso}"),
        )

        LogEntry.objects.log_action( user_id=self.request.user.pk,
            content_type_id=ContentType.objects.get_for_model(grupo_acesso).pk,
            object_id=grupo_acesso.pk,
            object_repr=force_str(grupo_acesso),
            action_flag=ADDITION,
            change_message=(f"Solicitado acesso aos recursos do Grupo de Trabalho por {colaborador.full_name} "),
        )

def add_header(canvas, doc, title):
    canvas.saveState()
    canvas.setFont('Helvetica-Bold', 12)

    header_text = title  
    image_path = "static/image/logo-coids.png"
    text_width = canvas.stringWidth(header_text, 'Helvetica-Bold', 12)
    x_position = (letter[0] - text_width) / 2

    image_width = 2 * cm
    image_height = 2 * cm    
    canvas.drawImage(image_path, 1 * cm, letter[1] - 2.5 * cm, width=image_width, height=image_height, preserveAspectRatio=True)

    canvas.drawString(x_position, letter[1] - 1.5 * cm, header_text)

    canvas.restoreState()


def export_to_xlsx(queryset, fields, title="Relatório", filename="relatorio.xlsx"):
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Dados"

        blue_fill = PatternFill(start_color="0077BA", end_color="0077BA", fill_type="solid")
        white_font = Font(name="Arial", color="FFFFFF", bold=True)
        thin_border = Border(bottom=Side(style='thin', color='000000'))

        title_font = Font(name='Arial', size=14, bold=True)
        header_font = Font(name='Arial', size=12, bold=True)
        data_font = Font(name='Arial', size=10)

        # Suporta dict de campos: {"campo": "Título"}
        if isinstance(fields, dict):
            field_names = list(fields.keys())
            header = list(fields.values())
        else:
            field_names = fields
            model = queryset.model
            header = [model._meta.get_field(f).verbose_name.capitalize() for f in fields]

        last_col = get_column_letter(len(field_names))
        ws.merge_cells(f'A1:{last_col}1')
        ws.row_dimensions[1].height = 50

        try:
            img_path = os.path.join(settings.BASE_DIR, 'static/image/logo-coids.png')
            pil_img = PILImage.open(img_path)
            pil_img.thumbnail((240, 60))
            img_io = BytesIO()
            pil_img.save(img_io, format='PNG')
            img_io.seek(0)
            excel_img = Image(img_io)
            ws.add_image(excel_img, 'A1')
        except Exception as e:
            print(f"Erro ao adicionar logo: {e}")
            ws['A1'] = "LOGO"
            ws['A1'].font = title_font

        ws['A1'].value = title
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_num, label in enumerate(header, 1):
            cell = ws.cell(row=2, column=col_num, value=label)
            cell.font = white_font
            cell.fill = blue_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border

        for row_num, obj in enumerate(queryset, 3):
            max_lines_in_row = 1
            for col_num, field in enumerate(field_names, 1):
                value = getattr(obj, field, "")
                cell_value = str(value) if value not in [None, ""] else "-"
                line_breaks = cell_value.count('\n') + 1
                if len(cell_value) > 100:
                    wrapped_lines = (len(cell_value) // 100) + 1
                    line_breaks = max(line_breaks, wrapped_lines)
                max_lines_in_row = max(max_lines_in_row, line_breaks)

            ws.row_dimensions[row_num].height = 15 * max_lines_in_row

            for col_num, field in enumerate(field_names, 1):
                value = getattr(obj, field, "")
                cell_value = str(value) if value not in [None, ""] else "-"
                cell = ws.cell(row=row_num, column=col_num, value=cell_value)
                cell.font = data_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        max_width = 80
        min_width = 20
        for column in ws.columns:
            col_letter = get_column_letter(column[0].column)
            field_name = field_names[column[0].column - 1] if column[0].column <= len(field_names) else None

            if field_name == 'grupo':
                max_length = 0
                for cell in column:
                    if cell.value and cell.value != '-':
                        groups = str(cell.value).split('\n') if '\n' in str(cell.value) else [str(cell.value)]
                        max_group_length = max(len(group.strip()) for group in groups)
                        max_length = max(max_length, max_group_length)
                adjusted_width = min(max(20, max_length * 1.1), 40)
                ws.column_dimensions[col_letter].width = adjusted_width
                continue

            max_length = max(
                (len(str(cell.value or "")) for cell in column),
                default=min_width
            )
            adjusted_width = min(max(min_width, max_length * 1.1), max_width)
            ws.column_dimensions[col_letter].width = adjusted_width

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{filename}"'},
        )
        wb.save(response)
        return response

    except Exception as err:
        print(f'Erro ao gerar Excel: {str(err)}')
        return HttpResponse(f"Erro ao gerar arquivo: {str(err)}", status=500)


def export_to_pdf(queryset=None, fields=None, filename=None, page_title=None):
    try:
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename={filename}'
        title = page_title or "Dados exportados para PDF"
        doc = SimpleDocTemplate(response, pagesize=letter)

        # Converte lista de campos em dicionário com verbose_name (se possível)
        if isinstance(fields, dict):
            field_map = fields
        else:
            model = queryset.model
            field_map = {}
            for f in fields:
                try:
                    verbose = model._meta.get_field(f).verbose_name.capitalize()
                except:
                    verbose = f.capitalize()
                field_map[f] = verbose

        field_names = list(field_map.keys())
        headers = list(field_map.values())
        data = [headers]

        for obj in queryset:
            row = []
            for field in field_names:
                try:
                    value = getattr(obj, field, '-')
                    value = value() if callable(value) else value
                    row.append(str(value) if value not in [None, ""] else '-')
                except Exception as e:
                    print(f"[ERRO] Campo '{field}' falhou em {obj}: {e}")
                    row.append('-')
            data.append(row)

        table = Table(data)

        style = TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), '#0077BA'),
            ('TEXTCOLOR', (0, 0), (-1, 0), (1, 1, 1)),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.5, (0, 0, 0)),
        ])
        table.setStyle(style)

        doc.build([table],
                  onFirstPage=lambda c, d: add_header(c, d, title),
                  onLaterPages=lambda c, d: add_header(c, d, title))

        return response
    except Exception as err:
        print(f'CARA DE ERRO AQUI VEY: {err}')
        return None


